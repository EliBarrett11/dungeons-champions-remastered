#!/usr/bin/env python3
"""Roster importer (M3): converts the approved roster workbook into content/roster/champions.json.

The workbook in content/roster/source/ is the source record; champions.json is the single copy the
game reads. Run this script after the workbook changes, then run the checks (see docs/PROGRESS.md).

  python3 -I tools/import_roster.py           # regenerate content/roster/champions.json
  python3 -I tools/import_roster.py --check   # exit 1 if champions.json differs from the workbook

It copies the workbook's text and numbers exactly. The only transformation is D-020: Feat Milestone
2 text that describes a Multi-Class Feat is left out of the playable data and recorded under
`excluded`. It never edits Class, Race or Origin assignments; roster rules (135 Champions, unique
combinations, 9 per trait) are checked by `npm run check` (apps/web roster validator) and the JSON
Schema by `npm run content-check`.

Requires Python 3 and openpyxl (exit code 2 if openpyxl is missing).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - depends on the environment
    print("import-roster: cannot run: the Python 'openpyxl' package is required", file=sys.stderr)
    sys.exit(2)

REPO = Path(__file__).resolve().parent.parent
WORKBOOK = Path("content/roster/source/Dungeons_and_Champions_Roster_135_Option_A_Unique_Combinations.xlsx")
OUTPUT = Path("content/roster/champions.json")
SHEET = "Champion Roster"

# Exact header row expected in the workbook, mapped to the JSON field it fills.
COLUMNS = [
    ("Roster ID", "id"),
    ("Champion Name (PROPOSED)", "name"),
    ("Cost Rank (PROPOSED)", "proposed.cost_rank"),
    ("Rarity (PROPOSED)", "proposed.rarity"),
    ("Role (PROPOSED)", "proposed.role"),
    ("Primary Class", "class"),
    ("Race", "race"),
    ("Origin", "origin"),
    ("Basic Attack (PROPOSED)", "proposed.basic_attack"),
    ("Passive (PROPOSED)", "proposed.passive"),
    ("Ultimate (PROPOSED)", "proposed.ultimate"),
    ("Feat Milestone 1 (PROPOSED)", "proposed.feat_milestone_1"),
    ("Feat Milestone 2 (PROPOSED)", "proposed.feat_milestone_2"),
    ("Champion Status", "proposed.design_status"),
    ("Notes", "proposed.notes"),
]

APPROVAL = (
    "Roster IDs, names and Class/Race/Origin assignments: approved by the owner on 2026-10-11 "
    "(DECISIONS.md D-017 to D-021). Everything under `proposed` (cost, rarity, Role, ability and Feat "
    "text, design status, notes) is PROPOSED content: preserved from the workbook, not confirmed, "
    "and not executed by the engine."
)

MULTI_CLASS_FEAT = re.compile(r"secondary .* Trait|multi-?class", re.IGNORECASE)
EXCLUDED_REASON = "Multi-Class Feat text excluded (D-020)"


class ImportError_(Exception):
    pass


def read_workbook(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if SHEET not in workbook.sheetnames:
        raise ImportError_(f"sheet {SHEET!r} not found; sheets are {workbook.sheetnames}")
    rows = list(workbook[SHEET].iter_rows(values_only=True))
    header = [cell for cell in rows[0]]
    expected = [title for title, _ in COLUMNS]
    if header != expected:
        raise ImportError_(f"unexpected header row:\n  found    {header}\n  expected {expected}")

    champions = []
    for number, row in enumerate(rows[1:], start=2):
        if all(cell is None for cell in row):
            continue
        record: dict = {"proposed": {}}
        for (title, field), value in zip(COLUMNS, row):
            if isinstance(value, str) and value.startswith("="):
                raise ImportError_(f"row {number}, {title!r}: formulas are not supported ({value!r})")
            if field == "proposed.cost_rank":
                if not isinstance(value, int):
                    raise ImportError_(f"row {number}, {title!r}: expected a whole number, found {value!r}")
            elif not isinstance(value, str):
                raise ImportError_(f"row {number}, {title!r}: expected text, found {value!r}")
            if field.startswith("proposed."):
                record["proposed"][field.split(".", 1)[1]] = value
            else:
                record[field] = value
        feat_2 = record["proposed"].get("feat_milestone_2", "")
        excluded = None
        if MULTI_CLASS_FEAT.search(feat_2):
            del record["proposed"]["feat_milestone_2"]
            excluded = {"feat_milestone_2": EXCLUDED_REASON}
        ordered = {key: record[key] for key in ("id", "name", "class", "race", "origin", "proposed")}
        if excluded:
            ordered["excluded"] = excluded
        champions.append(ordered)
    return champions


def build_document(root: Path) -> str:
    source = root / WORKBOOK
    if not source.exists():
        raise ImportError_(f"workbook not found: {WORKBOOK}")
    document = {
        "format": "dc-roster-v1",
        "approval": APPROVAL,
        "source": {
            "file": WORKBOOK.as_posix(),
            "sheet": SHEET,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        },
        "champions": read_workbook(source),
    }
    return json.dumps(document, indent=2, ensure_ascii=False) + "\n"


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Import the approved roster workbook (M3).")
    parser.add_argument("--check", action="store_true", help="only verify that champions.json matches the workbook")
    parser.add_argument("--root", type=Path, default=REPO, help="repository root")
    args = parser.parse_args(argv)
    root: Path = args.root.resolve()

    try:
        text = build_document(root)
    except ImportError_ as error:
        print(f"import-roster: FAIL: {error}")
        return 1

    output = root / OUTPUT
    count = json.loads(text)["champions"]
    excluded = sum(1 for champion in count if "excluded" in champion)
    if args.check:
        current = output.read_text(encoding="utf-8") if output.exists() else None
        if current != text:
            print(f"import-roster: FAIL: {OUTPUT} does not match {WORKBOOK}; run tools/import_roster.py")
            return 1
        print(f"import-roster: PASS: {OUTPUT} matches the workbook ({len(count)} Champions, {excluded} Multi-Class Feat texts excluded)")
        return 0
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(text, encoding="utf-8")
    print(f"import-roster: wrote {OUTPUT} ({len(count)} Champions, {excluded} Multi-Class Feat texts excluded)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
